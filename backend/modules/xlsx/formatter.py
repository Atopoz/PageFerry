"""把确认安全的 XLSX 译文写回 cell, 并同步 Excel Table 表头。"""

import re
from collections import defaultdict
from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import load_workbook

from .entities import XlsxCellSegment


class XlsxFormatter:
    """只替换提取阶段确认过的字符串 cell, 不重建 workbook 结构。"""

    def apply(
        self,
        source_path: str | Path,
        output_path: str | Path,
        segments: Sequence[XlsxCellSegment],
        translations: Mapping[tuple[int, str], str],
    ) -> None:
        """写回译文, 并在 Table header 重名时保留原 metadata。"""

        workbook = load_workbook(source_path, data_only=False)
        table_updates: dict[tuple[int, str], dict[int, str]] = defaultdict(dict)
        try:
            for segment in segments:
                translated = translations.get(segment.key, segment.original_text)
                worksheet = workbook.worksheets[segment.sheet_index - 1]
                cell = worksheet[segment.cell_ref]
                # 即使提取和写回之间对象发生异常变化, 也绝不覆盖公式。
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    continue
                cell.value = translated
                for header_ref in segment.table_headers:
                    table_updates[(segment.sheet_index, header_ref.table_name)][
                        header_ref.column_index
                    ] = translated
            self._apply_table_header_updates(workbook, table_updates)
            workbook.save(output_path)
        finally:
            workbook.close()

    def _apply_table_header_updates(
        self,
        workbook: object,
        updates: Mapping[tuple[int, str], Mapping[int, str]],
    ) -> None:
        """只提交不会产生重复列名的 Table header metadata。"""

        for (sheet_index, table_name), column_updates in updates.items():
            worksheet = workbook.worksheets[sheet_index - 1]
            table = worksheet.tables.get(table_name)
            if table is None:
                continue
            table_columns = list(getattr(table, "tableColumns", ()) or ())
            current_names = [str(column.name or "") for column in table_columns]
            for column_index in sorted(column_updates):
                if not 1 <= column_index <= len(table_columns):
                    continue
                normalized_name = self._normalize_table_header(column_updates[column_index])
                if not normalized_name:
                    continue
                candidate_names = current_names.copy()
                candidate_names[column_index - 1] = normalized_name
                if len(set(candidate_names)) != len(candidate_names):
                    continue
                table_columns[column_index - 1].name = normalized_name
                current_names[column_index - 1] = normalized_name

    @staticmethod
    def _normalize_table_header(text: str) -> str:
        """把 Table metadata 约束为单行非空名称。"""

        normalized = text.replace("\r", " ").replace("\n", " ").strip()
        return re.sub(r"\s+", " ", normalized)
